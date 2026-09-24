import os
import time
import datetime
from decimal import Decimal
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone

from apps.vehicles.models import DataSource


class CleanWorksheet:
    """Wrapper around openpyxl worksheet that automatically skips completely empty rows."""
    def __init__(self, ws):
        self._ws = ws

    def iter_rows(self, *args, **kwargs):
        for row in self._ws.iter_rows(*args, **kwargs):
            if row and any(c is not None and str(c).strip() != "" for c in row):
                yield row

    def __getattr__(self, name):
        return getattr(self._ws, name)


class BaseSheetSyncCommand(BaseCommand):
    """
    Abstract Base Class for Sheet-Level Data Synchronization Commands.
    Enforces idempotency, integer-cent monetary precision, atomic transactions,
    clean error handling, and audit logging to DataSource.
    """
    default_workbook_path: str = ""
    sheet_name: str = ""
    source_code: str = ""
    source_name: str = ""

    def add_arguments(self, parser):
        parser.add_argument(
            "--filepath",
            type=str,
            default=self.default_workbook_path,
            help=f"Path to Excel workbook (default: {self.default_workbook_path})",
        )
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Simulate ingestion without writing changes to the database",
        )
        parser.add_argument(
            "--clear",
            action="store_true",
            help="Clear existing table records before syncing (if supported)",
        )

    # -------------------------------------------------------------------------
    # Safe Type Conversion Utilities
    # -------------------------------------------------------------------------
    def _safe_str(self, val: Any) -> str:
        if val is None:
            return ""
        s = str(val).strip()
        return "" if s.lower() == "nan" else s

    def _safe_float(self, val: Any) -> Optional[float]:
        if val is None or str(val).strip() in ("", "nan", "None", "-"):
            return None
        try:
            return float(val)
        except (ValueError, TypeError):
            return None

    def _safe_int(self, val: Any) -> Optional[int]:
        if val is None or str(val).strip() in ("", "nan", "None", "-"):
            return None
        try:
            return int(float(val))
        except (ValueError, TypeError):
            return None

    def _safe_bool(self, val: Any) -> bool:
        if val is None:
            return False
        s = str(val).strip().lower()
        return s in ("yes", "true", "1", "t", "y", "active")

    def _dollars_to_cents(self, val: Any) -> Optional[int]:
        """Convert a dollar value to whole integer cents ($100.00 -> 10000)."""
        f = self._safe_float(val)
        if f is None:
            return None
        return int(round(f * 100))

    def _parse_date(self, val: Any) -> Optional[datetime.date]:
        if val is None or str(val).strip() in ("", "nan", "None", "-"):
            return None
        if isinstance(val, (datetime.date, datetime.datetime)):
            return val.date() if isinstance(val, datetime.datetime) else val
        s = str(val).strip()[:10]
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%Y/%m/%d"):
            try:
                return datetime.datetime.strptime(s, fmt).date()
            except ValueError:
                continue
        return None

    def _row_val(self, row: Any, hmap: Dict[str, int], col_name: str, default: Any = None) -> Any:
        """Safely fetch cell value by column name, returning default if column does not exist or index out of bounds."""
        if not hmap or col_name not in hmap:
            return default
        idx = hmap[col_name]
        if idx is None or idx >= len(row):
            return default
        val = row[idx]
        return val if val is not None else default

    # -------------------------------------------------------------------------
    # Excel Loading Helper
    # -------------------------------------------------------------------------
    def load_sheet(self, filepath: str, sheet_name: str):
        resolved_path = os.path.abspath(filepath)
        if not os.path.exists(resolved_path):
            raise CommandError(f"Workbook file not found at: {resolved_path}")

        self.stdout.write(f"Loading '{sheet_name}' from {resolved_path}...")
        try:
            wb = openpyxl.load_workbook(resolved_path, read_only=True, data_only=True)
            if sheet_name not in wb.sheetnames:
                raise CommandError(
                    f"Sheet '{sheet_name}' not found. Available sheets: {wb.sheetnames}"
                )
            return CleanWorksheet(wb[sheet_name])
        except Exception as e:
            raise CommandError(f"Failed to load Excel workbook: {e}")

    # -------------------------------------------------------------------------
    # Audit Logging to DataSource
    # -------------------------------------------------------------------------
    def log_sync_audit(
        self,
        created: int,
        updated: int,
        errors: List[str],
        duration_sec: float,
        is_dry_run: bool,
    ):
        if is_dry_run or not self.source_code:
            return

        status = "SUCCESS" if not errors else ("PARTIAL" if (created + updated) > 0 else "FAILED")
        now = timezone.now()
        DataSource.objects.update_or_create(
            code=self.source_code,
            defaults={
                "name": self.source_name or self.source_code,
                "dataset_type": "xlsx",
                "retrieved_at": now,
                "is_active": True,
                "metadata": {
                    "last_sync_at": now.isoformat(),
                    "sheet_name": self.sheet_name,
                    "rows_created": created,
                    "rows_updated": updated,
                    "duration_seconds": round(duration_sec, 2),
                    "status": status,
                    "errors": errors[:50],  # cap at 50 errors
                },
            },
        )

    # -------------------------------------------------------------------------
    # Handle & Execution Flow
    # -------------------------------------------------------------------------
    def handle(self, *args, **options):
        filepath = options["filepath"] or self.default_workbook_path
        is_dry_run = options["dry_run"]
        clear = options["clear"]

        if is_dry_run:
            self.stdout.write(self.style.WARNING("Running in --dry-run mode. No changes will be saved."))

        start_time = time.time()
        sheet = self.load_sheet(filepath, self.sheet_name)

        created, updated, errors = 0, 0, []
        try:
            if is_dry_run:
                created, updated, errors = self.sync_sheet(sheet, is_dry_run=True, clear=clear)
            else:
                with transaction.atomic():
                    created, updated, errors = self.sync_sheet(sheet, is_dry_run=False, clear=clear)
        except Exception as e:
            errors.append(str(e))
            self.stderr.write(self.style.ERROR(f"Sync failed with exception: {e}"))
            raise

        duration = time.time() - start_time
        self.log_sync_audit(created, updated, errors, duration, is_dry_run)

        self.stdout.write(
            self.style.SUCCESS(
                f"[{self.sheet_name}] Finished in {duration:.2f}s | "
                f"Created: {created} | Updated: {updated} | Errors: {len(errors)}"
            )
        )

    def sync_sheet(self, sheet, is_dry_run: bool = False, clear: bool = False) -> Tuple[int, int, List[str]]:
        """Subclasses implement specific row-by-row parsing and upsert logic."""
        raise NotImplementedError("Subclasses must implement sync_sheet()")
