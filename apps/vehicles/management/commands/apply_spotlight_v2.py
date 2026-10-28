import os

import openpyxl
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from apps.vehicles.models import Vehicle

BADGE_FLAGS = {
    "range kings": "range_kings",
    "green & lean": "green_lean",
    "unicorn match": "unicorn_match",
    "core all-stars": "core_allstars",
    "category sweep": "category_sweep",
}

VERSION = "v6.4-spotlight-rules-v2 (27-Oct-26)"


class Command(BaseCommand):
    help = (
        "Apply the Spotlight rules v2 changeset (27-Oct-26) to Vehicle.spotlight_data.\n\n"
        "Reads the 'Changes' sheet (Make/Model/Variant/Year/Badges after) and rewrites\n"
        "the badge flags/count/badges for each matched vehicle."
    )

    def add_arguments(self, parser):
        parser.add_argument("filepath", type=str, help="Path to 'Spotlight rules v2 changes 27-Oct-26.xlsx'")
        parser.add_argument("--dry-run", action="store_true", help="Report matches only; no writes")

    def handle(self, *args, **options):
        filepath = options["filepath"]
        dry_run = options["dry_run"]

        if not os.path.exists(filepath):
            raise CommandError(f"File not found: {filepath}")

        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        if "Changes" not in wb.sheetnames:
            raise CommandError(f"'Changes' sheet not found. Available: {wb.sheetnames}")
        ws = wb["Changes"]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()

        if len(rows) < 2:
            raise CommandError("Changes sheet has no data rows.")

        headers = [str(c).strip().lower() if c else "" for c in rows[0]]
        hmap = {h: i for i, h in enumerate(headers) if h}
        for col in ("make", "model", "variant", "year", "badges after"):
            if col not in hmap:
                raise CommandError(f"Missing expected column '{col}' in Changes sheet.")

        applied, not_found, ambiguous = 0, [], []

        def parse_badges(raw):
            flags = {v: False for v in BADGE_FLAGS.values()}
            badges = []
            if raw:
                for part in str(raw).split(","):
                    name = part.replace("™", "").strip()
                    if not name:
                        continue
                    badges.append(part.strip())
                    key = BADGE_FLAGS.get(name.lower())
                    if key:
                        flags[key] = True
            return flags, badges

        with transaction.atomic():
            for row in rows[1:]:
                make = str(row[hmap["make"]]).strip() if row[hmap["make"]] else ""
                model = str(row[hmap["model"]]).strip() if row[hmap["model"]] else ""
                variant = str(row[hmap["variant"]]).strip() if row[hmap["variant"]] else ""
                year_raw = row[hmap["year"]]
                fcai = str(row[hmap["fcai_category"]]).strip() if "fcai_category" in hmap and row[hmap["fcai_category"]] else ""
                badges_after = row[hmap["badges after"]]

                if not make or not model:
                    continue
                try:
                    year = int(float(year_raw))
                except (TypeError, ValueError):
                    not_found.append(f"{make} {model} {variant} (bad year: {year_raw})")
                    continue

                qs = Vehicle.objects.filter(
                    brand__iexact=make, model__iexact=model,
                    variant__iexact=variant, year=year,
                )
                if qs.count() > 1 and fcai:
                    qs = qs.filter(fcai_category__iexact=fcai)
                count = qs.count()
                if count == 0:
                    not_found.append(f"{make} {model} {variant} {year}")
                    continue
                if count > 1:
                    ambiguous.append(f"{make} {model} {variant} {year} ({count} matches)")
                    continue

                flags, badges = parse_badges(badges_after)
                vehicle = qs.first()
                data = dict(vehicle.spotlight_data or {})
                data.update(flags)
                data["count"] = len(badges)
                data["badges"] = ", ".join(badges)
                data["version"] = VERSION
                if not dry_run:
                    vehicle.spotlight_data = data
                    vehicle.save(update_fields=["spotlight_data"])
                applied += 1

        self.stdout.write(f"Rows in changeset: {len(rows) - 1}")
        self.stdout.write(self.style.SUCCESS(f"Applied: {applied}"))
        if ambiguous:
            self.stdout.write(self.style.WARNING(f"Ambiguous (skipped): {len(ambiguous)}"))
            for a in ambiguous:
                self.stdout.write(f"  - {a}")
        if not_found:
            self.stdout.write(self.style.WARNING(f"Not found (skipped): {len(not_found)}"))
            for n in not_found:
                self.stdout.write(f"  - {n}")
        if dry_run:
            self.stdout.write(self.style.WARNING("Dry run — nothing was written."))
