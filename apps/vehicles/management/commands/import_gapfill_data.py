import os
import datetime
import openpyxl
from decimal import Decimal
from typing import Any, Dict, Optional, Tuple

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from apps.vehicles.models import (
    DataSource,
    VehicleIdentity,
    ManualVehicleEnrichment,
    PowertrainType,
    TestCycle,
)


class Command(BaseCommand):
    help = "Safely ingest vehicle records from 'Carreb - Market Gap-Fill Dataset TEMPLATE v1.0.xlsx' into VehicleIdentity and DataSource models."

    POWERTRAIN_MAP = {
        "petrol": PowertrainType.PETROL,
        "diesel": PowertrainType.DIESEL,
        "hybrid": PowertrainType.HYBRID,
        "plug-in hybrid": PowertrainType.PHEV,
        "phev": PowertrainType.PHEV,
        "electric": PowertrainType.BEV,
        "bev": PowertrainType.BEV,
        "hydrogen": PowertrainType.FCEV,
        "fcev": PowertrainType.FCEV,
    }

    def add_arguments(self, parser):
        parser.add_argument("filepath", type=str, help="Path to the Gap-Fill Excel (.xlsx) file")
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Simulate import and validate rows without committing database changes",
        )
        parser.add_argument(
            "--include-example",
            action="store_true",
            help="Include row 3 (worked example) during ingestion (default is skipped)",
        )

    def _safe_str(self, val: Any) -> str:
        if val is None:
            return ""
        return str(val).strip()

    def _safe_float(self, val: Any) -> Optional[float]:
        if val is None or str(val).strip() == "":
            return None
        try:
            return float(val)
        except (ValueError, TypeError):
            return None

    def _safe_int(self, val: Any) -> Optional[int]:
        if val is None or str(val).strip() == "":
            return None
        try:
            return int(float(val))
        except (ValueError, TypeError):
            return None

    def _dollars_to_cents(self, val: Any) -> Optional[int]:
        f = self._safe_float(val)
        if f is None:
            return None
        return int(round(f * 100))

    def _parse_bool_on_sale(self, val: Any) -> Tuple[Optional[bool], str]:
        s = self._safe_str(val).lower()
        if s in ("yes", "runout"):
            return True, s
        if s == "announced":
            return False, s
        return None, s

    def _parse_date(self, val: Any) -> Optional[datetime.date]:
        if val is None or str(val).strip() == "":
            return None
        if isinstance(val, (datetime.date, datetime.datetime)):
            return val.date() if isinstance(val, datetime.datetime) else val
        try:
            return datetime.datetime.strptime(str(val).strip()[:10], "%Y-%m-%d").date()
        except ValueError:
            return None

    def _map_powertrain(self, fuel_type_raw: str) -> str:
        clean = fuel_type_raw.lower().strip()
        return self.POWERTRAIN_MAP.get(clean, PowertrainType.PETROL)

    def handle(self, *args, **options):
        filepath = options["filepath"]
        is_dry_run = options["dry_run"]
        include_example = options["include_example"]

        if not os.path.exists(filepath):
            raise CommandError(f"Excel file not found at: {filepath}")

        self.stdout.write(self.style.NOTICE(f"Loading workbook: {filepath}"))
        wb = openpyxl.load_workbook(filepath, data_only=True)

        if "Data" not in wb.sheetnames:
            raise CommandError("Missing required sheet 'Data' in workbook.")

        # 1. Ingest DataSources from 'Brand tracker' if present
        data_sources_by_brand: Dict[str, DataSource] = {}
        if "Brand tracker" in wb.sheetnames:
            self.stdout.write("Processing 'Brand tracker' sheet for provenance tracking...")
            ws_bt = wb["Brand tracker"]
            for r in range(4, ws_bt.max_row + 1):
                brand = self._safe_str(ws_bt.cell(r, 1).value)
                if not brand or "total" in brand.lower() or "example" in brand.lower():
                    continue

                code = f"GAPFILL_{brand.upper()[:20]}_{datetime.date.today().year}"
                source_name = f"Market Gap-Fill — {brand}"
                primary_src = self._safe_str(ws_bt.cell(r, 4).value)
                date_completed = self._parse_date(ws_bt.cell(r, 5).value)
                notes = self._safe_str(ws_bt.cell(r, 6).value)

                metadata = {
                    "models_found": self._safe_int(ws_bt.cell(r, 2).value),
                    "variants_recorded": self._safe_int(ws_bt.cell(r, 3).value),
                    "primary_source_used": primary_src,
                    "notes": notes,
                }

                if not is_dry_run:
                    ds, _ = DataSource.objects.update_or_create(
                        code=code,
                        defaults={
                            "name": source_name,
                            "url": primary_src if primary_src.startswith("http") else "",
                            "dataset_type": "gap_fill",
                            "retrieved_at": (
                                datetime.datetime.combine(date_completed, datetime.time.min)
                                if date_completed
                                else datetime.datetime.now()
                            ),
                            "metadata": metadata,
                            "is_active": True,
                        },
                    )
                    data_sources_by_brand[brand.lower()] = ds
                else:
                    self.stdout.write(f"  [DRY-RUN] Discovered DataSource: {code} ({source_name})")

        # 2. Ingest Vehicle Data rows from 'Data' sheet
        ws_data = wb["Data"]
        header_cols = [self._safe_str(ws_data.cell(2, c).value) for c in range(1, ws_data.max_column + 1)]
        col_idx = {name: idx + 1 for idx, name in enumerate(header_cols) if name}

        start_row = 3 if include_example else 4
        self.stdout.write(f"Processing 'Data' sheet from row {start_row} to {ws_data.max_row}...")

        created_count = 0
        updated_count = 0
        skipped_count = 0

        with transaction.atomic():
            for r in range(start_row, ws_data.max_row + 1):
                make = self._safe_str(ws_data.cell(r, col_idx.get("make", 1)).value)
                model = self._safe_str(ws_data.cell(r, col_idx.get("model", 2)).value)
                variant = self._safe_str(ws_data.cell(r, col_idx.get("variant", 3)).value)
                year = self._safe_int(ws_data.cell(r, col_idx.get("model_year", 4)).value)

                if not make or not model or not year:
                    skipped_count += 1
                    continue

                fuel_type_raw = self._safe_str(ws_data.cell(r, col_idx.get("fuel_type", 6)).value)
                powertrain = self._map_powertrain(fuel_type_raw)
                is_current, sales_status = self._parse_bool_on_sale(
                    ws_data.cell(r, col_idx.get("on_sale", 5)).value
                )

                # Pricing & Specs
                price_aud = ws_data.cell(r, col_idx.get("price", 20)).value
                driveaway_cents = self._dollars_to_cents(price_aud)

                # Packaging raw data dictionary for full auditability
                raw_dict = {}
                for col_name, c_num in col_idx.items():
                    raw_dict[col_name] = ws_data.cell(r, c_num).value

                # Provenance dictionary
                provenance = {
                    "source_name": self._safe_str(ws_data.cell(r, col_idx.get("source_name", 24)).value),
                    "source_url": self._safe_str(ws_data.cell(r, col_idx.get("source_url", 25)).value),
                    "retrieved_at": self._safe_str(ws_data.cell(r, col_idx.get("retrieved_at", 26)).value),
                    "price_type": self._safe_str(ws_data.cell(r, col_idx.get("price_type", 21)).value),
                    "price_state": self._safe_str(ws_data.cell(r, col_idx.get("price_state", 22)).value),
                    "price_as_at": self._safe_str(ws_data.cell(r, col_idx.get("price_as_at", 23)).value),
                    "sales_status": sales_status,
                    "drive": self._safe_str(ws_data.cell(r, col_idx.get("drive", 8)).value),
                    "tow_braked_kg": self._safe_int(ws_data.cell(r, col_idx.get("tow_braked_kg", 15)).value),
                    "tow_unbraked_kg": self._safe_int(ws_data.cell(r, col_idx.get("tow_unbraked_kg", 16)).value),
                    "ancap_tested_year": self._safe_int(ws_data.cell(r, col_idx.get("ancap_tested_year", 18)).value),
                }

                drive = self._safe_str(ws_data.cell(r, col_idx.get("drive", 8)).value)
                tow_braked = self._safe_int(ws_data.cell(r, col_idx.get("tow_braked_kg", 15)).value)
                tow_unbraked = self._safe_int(ws_data.cell(r, col_idx.get("tow_unbraked_kg", 16)).value)
                ancap_stars = self._safe_int(ws_data.cell(r, col_idx.get("ancap_stars", 17)).value)
                ancap_year = self._safe_int(ws_data.cell(r, col_idx.get("ancap_tested_year", 18)).value)
                price_type = self._safe_str(ws_data.cell(r, col_idx.get("price_type", 21)).value)
                price_state = self._safe_str(ws_data.cell(r, col_idx.get("price_state", 22)).value)
                price_as_at = self._parse_date(ws_data.cell(r, col_idx.get("price_as_at", 23)).value)
                notes = self._safe_str(ws_data.cell(r, col_idx.get("notes", 28)).value)

                ds = data_sources_by_brand.get(make.lower())

                vi_defaults = {
                    "source": ds,
                    "body_style": self._safe_str(ws_data.cell(r, col_idx.get("body_type", 13)).value),
                    "fcai_category": self._safe_str(ws_data.cell(r, col_idx.get("fcai_segment", 19)).value),
                    "powertrain_type": powertrain,
                    "fuel_type_raw": fuel_type_raw,
                    "engine_l": self._safe_float(ws_data.cell(r, col_idx.get("engine_l", 7)).value),
                    "battery_kwh": self._safe_float(ws_data.cell(r, col_idx.get("battery_kwh", 9)).value),
                    "electric_range_km": self._safe_float(ws_data.cell(r, col_idx.get("range_km", 10)).value),
                    "fuel_consumption_combined_l_100km": self._safe_float(
                        ws_data.cell(r, col_idx.get("fuel_l_100km", 11)).value
                    ),
                    "tailpipe_co2_g_per_km": self._safe_float(ws_data.cell(r, col_idx.get("co2_g_km", 12)).value),
                    "num_seats": self._safe_int(ws_data.cell(r, col_idx.get("seats", 14)).value),
                    "drive": drive,
                    "tow_braked_kg": tow_braked,
                    "tow_unbraked_kg": tow_unbraked,
                    "driveaway_price_estimate": driveaway_cents,
                    "price_type": price_type,
                    "price_state": price_state,
                    "price_as_at": price_as_at,
                    "is_current_model": is_current,
                    "sales_status": sales_status,
                    "test_cycle": self._safe_str(ws_data.cell(r, col_idx.get("test_cycle", 27)).value),
                    "source_dataset": "Gap-Fill v1.0",
                    "provenance": provenance,
                    "raw_data": raw_dict,
                }

                if is_dry_run:
                    self.stdout.write(f"  [DRY-RUN] Row {r}: {year} {make} {model} {variant} | Driveaway: {driveaway_cents} cents | ANCAP: {ancap_stars}")
                    created_count += 1
                else:
                    vi, created = VehicleIdentity.objects.update_or_create(
                        make=make,
                        model=model,
                        year=year,
                        variant=variant,
                        defaults=vi_defaults,
                    )
                    if created:
                        created_count += 1
                    else:
                        updated_count += 1

                    # Upsert enrichment record if ancap or notes are present
                    if ancap_stars is not None or ancap_year is not None or notes:
                        ManualVehicleEnrichment.objects.update_or_create(
                            record=vi,
                            defaults={
                                "ancap_safety_rating": ancap_stars,
                                "ancap_tested_year": ancap_year,
                                "notes": notes,
                                "driveaway_price_estimate": driveaway_cents,
                                "enriched_by": "Gap-Fill Ingestion Pipeline",
                            },
                        )

            if is_dry_run:
                transaction.set_rollback(True)

        self.stdout.write(
            self.style.SUCCESS(
                f"\nFinished Ingestion Summary:\n"
                f"  Created: {created_count}\n"
                f"  Updated: {updated_count}\n"
                f"  Skipped (Empty/Header): {skipped_count}\n"
                f"  Dry Run: {is_dry_run}"
            )
        )
