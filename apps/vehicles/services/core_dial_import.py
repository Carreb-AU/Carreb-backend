"""Parser that turns the "Consolidated Vehicles v6.4" workbook sheet into
CORE Dial rows (Mode B / Pure Cost only).

Sheet layout (verified against the 07-Oct-26, 13-Oct-26 and 21-Oct-26
workbooks):
  * rows sorted by (Make, Model, Variant, Year) — a "name group";
  * a name group holds one or more listings; each listing has exactly
    8 states × 4 modes = 32 rows;
  * within a name group, consecutive rows form state runs (state changes =>
    new run); each run holds N quads of 4 rows (modes B, E, C, F) — one quad
    per listing — and the quads of a given state are in listing order.

So: split each state run into quads at every 'B' row; quad k of every state
belongs to listing k. version_id is the catalogue order (first appearance,
1..N). model_id is the first 12 hex of SHA-256("year|make|model").

Columns are located by header NAME, not position: the 21-Oct-26 workbook
inserted new leading columns (listing_id at 0, Make at 1, ...), so any
positional mapping breaks silently on newer files. Columns that are absent
from a given workbook (e.g. price_basis in early files, is_suppressed
pre-21-Oct) resolve to None and are treated as empty / not suppressed.

National scores (D30): price_score / emissions_score are the medians of the
listing's per-state Mode B scores, rounded half-up to 6 dp, and are copied
onto every state row.
"""
import hashlib
import logging
import statistics
from dataclasses import dataclass, field
from typing import Dict, List, Optional

import openpyxl

from apps.vehicles.services.core_dial import AU_STATES

logger = logging.getLogger(__name__)

SHEET_NAME = 'Consolidated Vehicles v6.4'

# Research Status values that count as "on sale" (section 2)
ON_SALE_STATUSES = ('v6.3.4 seed', 'Confirmed')

# Header names for the scalar columns the parser needs.
COL_MAKE = 'Make'
COL_MODEL = 'Model'
COL_VARIANT = 'Variant'
COL_YEAR = 'Year'
COL_FUEL = 'Fuel Type'
COL_CO2_GKM = 'CO2 Tailpipe (g/km)'
COL_STATE = 'State'
COL_FCAI = 'fcai_category'
COL_DRIVEAWAY = 'Driveaway Price ($)'
COL_MODE = 'CORE Mode'
COL_MODE_COO = 'Mode COO p.a. ($)'
COL_CO2_KG = 'Annual CO2 (kg)'
COL_PRICE_SCORE = 'price_score'
COL_EMISSIONS_SCORE = 'emissions_score'
COL_SCORING_VERSION = 'scoring_version'
COL_RESEARCH_STATUS = 'Research Status'
COL_PRICE_BASIS = 'price_basis'  # optional — absent from early workbooks
COL_IS_SUPPRESSED = 'is_suppressed'  # optional — added in the 21-Oct-26 file

# Mode B parts of the COO: part key -> sheet column name.
PART_COLS = {
    'fuel': 'fuel_cost_pa',
    'electricity': 'electricity_cost_pa',
    'ctp': 'ctp_pa',
    'comprehensive': 'comprehensive_pa',
    'registration': 'registration_pa',
    'servicing': 'maintenance_pa',
    'tyres': 'tyre_cost_pa',
    'road_user_charge': 'ruc_pa',
}

PART_KEYS = tuple(PART_COLS.keys())


def _build_column_map(header) -> Dict[str, int]:
    """header row (tuple) -> {column name: 0-based index}. First wins."""
    columns: Dict[str, int] = {}
    for i, name in enumerate(header):
        if name is None:
            continue
        name = str(name).strip()
        if name and name not in columns:
            columns[name] = i
    return columns


def _cell(row, columns: Dict[str, int], name: str):
    """Row tuple value for a header name; None when the column is absent."""
    idx = columns.get(name)
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _truthy(value) -> bool:
    if value is None or value is True or value is False:
        return bool(value)
    return str(value).strip().lower() in ('true', 'yes', '1')


@dataclass
class ParsedListing:
    """One catalogue listing: a named variant with 8 states × 4 mode rows."""
    make: str
    model: str
    variant: str
    year: int
    research_status: str
    states: Dict[str, dict]  # state -> {'B': row, 'E': row, 'C': row, 'F': row}
    is_suppressed: bool = False


@dataclass
class DialRowData:
    """A CoreDialRow not yet written to the database."""
    version_id: int
    model_id: str
    make: str
    model: str
    variant: str
    year: int
    fcai_category: str
    state: str
    driveaway_cents: int
    coo_pa_cents: int
    annual_co2_kg: float
    co2_g_per_km: float
    price_score: float
    emissions_score: float
    co2_cap: Optional[float]
    parts_cents: Dict[str, int]
    price_basis: str = ''


@dataclass
class ParsedWorkbook:
    scoring_version: str
    as_at: Optional[str]
    columns: Dict[str, int] = field(default_factory=dict)
    listings: List[ParsedListing] = field(default_factory=list)
    problems: List[str] = field(default_factory=list)
    skipped_suppressed: int = 0

    def dial_rows(self):
        """Flatten to dial row data for on-sale, unsuppressed listings."""
        from apps.vehicles.services.core_dial import (
            co2_cap_for, dollars_to_cents, round_half_up,
        )

        rows: List[DialRowData] = []
        for version_id, listing in enumerate(self.listings, start=1):
            if listing.research_status not in ON_SALE_STATUSES:
                continue
            if listing.is_suppressed:
                self.skipped_suppressed += 1
                continue
            # National scores: median of per-state Mode B scores, 6 dp (D30)
            price_scores, emissions_scores = [], []
            for st in AU_STATES:
                b = listing.states[st]['B']
                score = _cell(b, self.columns, COL_PRICE_SCORE)
                if isinstance(score, (int, float)):
                    price_scores.append(score)
                score = _cell(b, self.columns, COL_EMISSIONS_SCORE)
                if isinstance(score, (int, float)):
                    emissions_scores.append(score)
            if not price_scores or not emissions_scores:
                self.problems.append(
                    f"{listing.make} {listing.model} {listing.variant} {listing.year}: "
                    "missing per-state scores; skipped"
                )
                continue
            nat_price = round_half_up(statistics.median(price_scores), 6)
            nat_emissions = round_half_up(statistics.median(emissions_scores), 6)

            act_b = listing.states['ACT']['B']
            co2_g_per_km = float(_cell(act_b, self.columns, COL_CO2_GKM) or 0)
            co2_cap = co2_cap_for(
                str(_cell(act_b, self.columns, COL_FUEL) or '').lower(), co2_g_per_km
            )
            model_id = hashlib.sha256(
                f"{listing.year}|{listing.make}|{listing.model}".encode('utf-8')
            ).hexdigest()[:12]

            for st in AU_STATES:
                b = listing.states[st]['B']
                parts_cents = {}
                for key, colname in PART_COLS.items():
                    raw = _cell(b, self.columns, colname)
                    parts_cents[key] = dollars_to_cents(raw) if isinstance(raw, (int, float)) else 0
                coo_cents = dollars_to_cents(_cell(b, self.columns, COL_MODE_COO))
                parts_sum = sum(parts_cents.values())
                if parts_sum != coo_cents:
                    self.problems.append(
                        f"{listing.make} {listing.model} {listing.variant} "
                        f"{listing.year} {st}: Mode B COO {coo_cents}c != sum of "
                        f"parts {parts_sum}c"
                    )
                rows.append(DialRowData(
                    version_id=version_id,
                    model_id=model_id,
                    make=listing.make,
                    model=listing.model,
                    variant=listing.variant,
                    year=listing.year,
                    fcai_category=str(_cell(b, self.columns, COL_FCAI) or ''),
                    state=st,
                    driveaway_cents=dollars_to_cents(_cell(b, self.columns, COL_DRIVEAWAY)),
                    coo_pa_cents=coo_cents,
                    annual_co2_kg=float(_cell(b, self.columns, COL_CO2_KG) or 0),
                    co2_g_per_km=co2_g_per_km,
                    price_score=nat_price,
                    emissions_score=nat_emissions,
                    co2_cap=co2_cap,
                    parts_cents=parts_cents,
                    price_basis=str(_cell(b, self.columns, COL_PRICE_BASIS) or ''),
                ))
        return rows


def parse_name_group(name, group_rows, problems, columns) -> List[ParsedListing]:
    """Split one name group's rows into listings (see module docstring)."""
    make, model, variant, year = name
    # Split into state runs; each run holds N quads (B,E,C,F), one per listing.
    per_state: Dict[str, list] = {}
    order: List[str] = []
    for i, row in enumerate(group_rows):
        if _cell(row, columns, COL_MODE) == 'B':
            state = _cell(row, columns, COL_STATE)
            if state not in per_state:
                order.append(state)
            per_state.setdefault(state, []).append(group_rows[i:i + 4])

    missing = [st for st in AU_STATES if st not in per_state]
    if missing:
        problems.append(f"{make} {model} {variant} {year}: missing states {missing}")
        return []

    quad_counts = {len(quads) for quads in per_state.values()}
    if len(quad_counts) != 1:
        problems.append(f"{make} {model} {variant} {year}: uneven state runs")
        return []
    n_listings = quad_counts.pop()

    listings = []
    for k in range(n_listings):
        states = {}
        ok = True
        for st in AU_STATES:
            quad = per_state[st][k]
            if len(quad) != 4 or [
                _cell(r, columns, COL_MODE) for r in quad
            ] != ['B', 'E', 'C', 'F']:
                problems.append(
                    f"{make} {model} {variant} {year} {st} listing {k + 1}: "
                    "malformed mode quad"
                )
                ok = False
                break
            states[st] = {_cell(r, columns, COL_MODE): r for r in quad}
        if not ok:
            continue
        research = {
            _cell(states[st]['B'], columns, COL_RESEARCH_STATUS) for st in AU_STATES
        }
        if len(research) != 1:
            problems.append(f"{make} {model} {variant} {year}: mixed research status")
            continue
        # is_suppressed is uniform across a listing's rows; the ACT Mode B
        # row is the canonical read (absent column => not suppressed).
        suppressed = _truthy(_cell(states['ACT']['B'], columns, COL_IS_SUPPRESSED))
        listings.append(ParsedListing(
            make=make, model=model, variant=variant, year=year,
            research_status=research.pop(),
            states=states,
            is_suppressed=suppressed,
        ))
    return listings


def parse_workbook(path, as_at: Optional[str] = None) -> ParsedWorkbook:
    """Parse the Consolidated Vehicles v6.4 sheet of the given workbook."""
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        ws = wb[SHEET_NAME]
        rows_iter = ws.iter_rows(values_only=True)
        header = next(rows_iter, ())
        columns = _build_column_map(header)
        name_groups: Dict[tuple, list] = {}
        group_order: List[tuple] = []
        scoring_versions = []
        for row in rows_iter:
            if _cell(row, columns, COL_MAKE) is None:
                continue
            name = (
                _cell(row, columns, COL_MAKE),
                _cell(row, columns, COL_MODEL),
                _cell(row, columns, COL_VARIANT),
                _cell(row, columns, COL_YEAR),
            )
            if name not in name_groups:
                name_groups[name] = []
                group_order.append(name)
            name_groups[name].append(row)
            sv = _cell(row, columns, COL_SCORING_VERSION)
            if sv and sv not in scoring_versions:
                scoring_versions.append(sv)
    finally:
        wb.close()

    problems: List[str] = []
    listings: List[ParsedListing] = []
    for name in group_order:
        listings.extend(parse_name_group(name, name_groups[name], problems, columns))

    if not scoring_versions:
        problems.append('No scoring_version found in sheet')
        scoring_version = 'unknown'
    else:
        scoring_version = scoring_versions[0]
        if len(scoring_versions) > 1:
            problems.append(
                f'Multiple scoring versions in sheet: {scoring_versions}; '
                f'using {scoring_version}'
            )

    return ParsedWorkbook(
        scoring_version=scoring_version,
        as_at=as_at,
        columns=columns,
        listings=listings,
        problems=problems,
    )
